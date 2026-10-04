import logging
import re
from pathlib import Path
from statistics import mean, stdev

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from ..eval import EvaluationResult
from .catalogue import ANSWER_COL, ID_COL, QUERY_COL, QUERY_TOPIC, QUERY_TYPE, VL_COL
from .workspace import EvaluationMetadata

logger = logging.getLogger(__name__)

_NUMERIC_FORMAT = "0.000000"

_MAX_WIDTH = 60
_MAX_TEXT_WIDTH = 80

_ACTUAL_OUTPUT_COL = "Actual LLM Output"
_FIXED_COLS = (
    ID_COL,
    VL_COL,
    QUERY_TYPE,
    QUERY_TOPIC,
    QUERY_COL,
    ANSWER_COL,
    _ACTUAL_OUTPUT_COL,
)
_BIG_TEXT_COLUMNS = [QUERY_COL, ANSWER_COL, _ACTUAL_OUTPUT_COL]


def _table_name(sheet_title: str, used: set[str]) -> str:
    name = re.sub(r"\W", "_", sheet_title)
    if not name or not name[0].isalpha():
        name = f"t_{name}"
    candidate, suffix = name, 2
    while candidate in used:
        candidate = f"{name}_{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


def _table_width(sheet: Worksheet, min_col: int, max_col: int) -> float:
    px = 0.0
    for idx in range(min_col, max_col + 1):
        width = sheet.column_dimensions[get_column_letter(col_idx=idx)].width
        px += width * 7 + 5
    return px * (2.54 / 96)


def _add_column_chart(
    sheet: Worksheet,
    title: str,
    min_col: int,
    max_col: int,
    max_row: int,
) -> None:
    if max_row < 2 or max_col < min_col + 1:
        return

    chart = BarChart()
    chart.title = title
    chart.style = 10

    data = Reference(sheet, min_col=min_col + 1, max_col=max_col, min_row=1, max_row=max_row)
    categories = Reference(sheet, min_col=min_col, max_col=min_col, min_row=2, max_row=max_row)
    chart.add_data(data=data, titles_from_data=True)
    chart.set_categories(categories)

    chart.legend.position = "t"
    chart.legend.overlay = False

    chart.width = int(_table_width(sheet=sheet, min_col=min_col, max_col=max_col))
    chart.height = round(chart.width / 2.0)

    sheet.add_chart(chart=chart, anchor=f"{get_column_letter(min_col)}{max_row + 2}")


def _format_as_table(
    sheet: Worksheet,
    used_table_names: set[str],
    min_col: int = 1,
    max_col: int | None = None,
    max_row: int | None = None,
    name: str | None = None,
) -> None:
    max_col = sheet.max_column if max_col is None else max_col
    max_row = sheet.max_row if max_row is None else max_row
    if max_row < 2 or max_col < min_col + 1:
        return

    ref = f"{get_column_letter(min_col)}1:{get_column_letter(max_col)}{max_row}"
    table = Table(
        displayName=_table_name(sheet_title=name or sheet.title, used=used_table_names),
        ref=ref,
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table=table)

    sheet.freeze_panes = "A2"

    for col_idx in range(min_col, max_col + 1):
        letter = get_column_letter(col_idx=col_idx)
        header = sheet.cell(row=1, column=col_idx).value
        is_wrap = header in _BIG_TEXT_COLUMNS
        if is_wrap:
            sheet.column_dimensions[letter].width = _MAX_TEXT_WIDTH
        else:
            longest = max(
                (len(str(c.value)) for c in sheet[letter] if c.value is not None), default=0
            )
            sheet.column_dimensions[letter].width = min(max(longest + 4, 10), _MAX_WIDTH)

        alignment = (
            Alignment(horizontal="left", vertical="top", wrap_text=True)
            if is_wrap
            else Alignment(horizontal="center", vertical="center")
        )

        for cell in sheet.iter_rows(min_row=2, max_row=max_row, min_col=col_idx, max_col=col_idx):
            cell[0].alignment = alignment


def _unique_sheet_name(base: str, used_names: set[str]) -> str:
    base = base[:31]
    if base not in used_names:
        return base

    suffix = 2
    while True:
        tail = f"_{suffix}"
        candidate = f"{base[: 31 - len(tail)]}{tail}"
        if candidate not in used_names:
            return candidate
        suffix += 1


def _as_number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        return float(value)

    if isinstance(value, str):
        try:
            return float(value)
        except ValueError:
            return None

    return None


def _clean_header(header: tuple) -> list[str]:
    seen: set[str] = set()
    clean = []
    for i, val in enumerate(header, start=1):
        name = str(val).strip() if val is not None else f"Column{i}"
        while name in seen:
            name = f"{name}_{i}"
        seen.add(name)
        clean.append(name)
    return clean


def _mean(values: list[float]) -> float | None:
    return mean(values) if values else None


def _stdev(values: list[float]) -> float | None:
    return stdev(values) if values and len(values) > 1 else None


def _write_overview(
    sheet: Worksheet,
    metrics: list[str],
    rows: list[tuple[str, dict[str, list[float]]]],
    used_table_names: set[str],
) -> None:
    tables = (("AVERAGES", _mean), ("STD_DEVIATION", _stdev))
    width = len(metrics) + 1

    for idx, (table_name, aggregate) in enumerate(tables):
        start_col = 1 + idx * (width + 1)

        sheet.cell(row=1, column=start_col, value="Experiment")
        for offset, metric in enumerate(metrics, start=1):
            sheet.cell(row=1, column=start_col + offset, value=metric)

        for row_idx, (exp, values) in enumerate(rows, start=2):
            sheet.cell(row=row_idx, column=start_col, value=exp)
            for offset, metric in enumerate(metrics, start=1):
                cell = sheet.cell(
                    row=row_idx, column=start_col + offset, value=aggregate(values.get(metric, []))
                )
                cell.number_format = _NUMERIC_FORMAT

        _format_as_table(
            sheet=sheet,
            used_table_names=used_table_names,
            min_col=start_col,
            max_col=start_col + len(metrics),
            max_row=len(rows) + 1,
            name=table_name,
        )
        _add_column_chart(
            sheet=sheet,
            title=table_name,
            min_col=start_col,
            max_col=start_col + len(metrics),
            max_row=len(rows) + 1,
        )
        logger.debug("Created a bar chart for sheet ['%s']", sheet.title)


def _merge_files(input_paths: list[Path], output_path: str, metrics: list[str]) -> None:
    merged = Workbook()
    merged.remove(merged.active)
    overview = merged.create_sheet(title="Overview", index=0)

    used_names: set[str] = {"Overview"}
    used_table_names: set[str] = set()
    overview_rows: list[tuple[str, dict[str, list[float]]]] = []

    for path in input_paths:
        src_wb = load_workbook(path, data_only=True)
        for sheet_name in src_wb.sheetnames:
            src_sheet = src_wb[sheet_name]
            new_sheet_name = _unique_sheet_name(sheet_name, used_names)
            used_names.add(new_sheet_name)
            dst_sheet = merged.create_sheet(title=new_sheet_name)

            rows = src_sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                continue

            clean_header = _clean_header(header)
            dst_sheet.append(clean_header)

            metric_columns = {idx: name for idx, name in enumerate(clean_header) if name in metrics}
            values: dict[str, list[float]] = {m: [] for m in metric_columns.values()}

            for row in rows:
                dst_sheet.append(row)
                for idx, metric in metric_columns.items():
                    number = _as_number(row[idx]) if idx < len(row) else None
                    if number is not None:
                        values[metric].append(number)

            _format_as_table(sheet=dst_sheet, used_table_names=used_table_names)
            overview_rows.append((new_sheet_name, values))

        logger.info("Processed ['%s']", path.stem)
        src_wb.close()

    _write_overview(
        sheet=overview, metrics=metrics, rows=overview_rows, used_table_names=used_table_names
    )
    logger.info("Created an overview sheet")

    merged.save(output_path)
    logger.info("Saved merged workbook to ['%s']", output_path)


def _get_file_paths_r(dir: Path) -> list[Path]:
    return sorted(p for p in dir.rglob("*_results.xlsx") if p.is_file())


def merge_experiment_results(eval_metadata: EvaluationMetadata, tgt: Path) -> None:
    assert eval_metadata.root_path is not None

    input_paths = _get_file_paths_r(eval_metadata.root_path / "experiments")
    if not input_paths:
        logger.warning("No experiment results were found. Aborting merge")
        return

    _merge_files(input_paths=input_paths, output_path=str(tgt), metrics=eval_metadata.metrics)


def export_eval_results(exp_id: str, results: list[EvaluationResult], output_path: Path) -> None:
    assert results is not None, "no evaluation results to export"

    metric_cols: list[str] = []
    for result in results:
        for metric in result.metrics:
            if metric.name not in metric_cols:
                metric_cols.append(metric.name)

    columns = list(_FIXED_COLS) + metric_cols

    def _sort_key(res: EvaluationResult) -> tuple[int, object]:
        case_id = res.exchange.case.id
        try:
            return (0, int(case_id))
        except (TypeError, ValueError):
            return (1, case_id)

    sorted_res = sorted(results, key=_sort_key)

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = exp_id

    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = Font(bold=True)

    for row_idx, res in enumerate(sorted_res, start=2):
        exch = res.exchange
        case = exch.case

        values = {
            ID_COL: case.id,
            VL_COL: case.vl,
            QUERY_TYPE: case.query_type,
            QUERY_TOPIC: case.query_topic,
            QUERY_COL: case.query,
            ANSWER_COL: case.exp_answer,
            _ACTUAL_OUTPUT_COL: exch.llm_response,
        }

        for metric in res.metrics:
            values[metric.name] = str(metric.score)

        for col_idx, col_name in enumerate(columns, start=1):
            ws.cell(row=row_idx, column=col_idx, value=values.get(col_name))

    wb.save(output_path)
